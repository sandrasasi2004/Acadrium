import os
import sys
import time
import pymupdf
import docx
import pptx
from fastapi.testclient import TestClient

# Ensure backend package can be imported
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app.main import app
from app.database.session import SessionLocal
from app.services.document_processor import (
    process_uploaded_resource,
    process_workspace_resource,
    search_resources_by_text,
    search_workspace_files_by_text
)
from app.models.user import User

client = TestClient(app)

def create_sample_pdf(filepath: str, text: str):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    doc.save(filepath)
    doc.close()

def create_sample_docx(filepath: str, text: str):
    doc = docx.Document()
    doc.add_paragraph(text)
    doc.save(filepath)

def create_sample_pptx(filepath: str, text: str):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = text
    prs.save(filepath)

def create_sample_txt(filepath: str, text: str):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(text)

def test_document_processing_suite():
    print("\n============================================================")
    print("STARTING DOCUMENT PROCESSING & TEXT EXTRACTION TEST SUITE")
    print("============================================================")

    timestamp = int(time.time())
    fac_email = f"fac_doc_{timestamp}@test.com"
    password = "Password123!"
    temp_files = []

    # 1. Register & Login Faculty User
    reg_resp = client.post("/api/auth/register", json={
        "username": "Dr. Document Processor",
        "email": fac_email,
        "password": password,
        "role": "faculty",
        "department": "Computer Science"
    })
    assert reg_resp.status_code == 200, f"Registration failed: {reg_resp.text}"

    login_resp = client.post("/api/auth/login", json={"email": fac_email, "password": password})
    assert login_resp.status_code == 200, f"Login failed: {login_resp.text}"
    token = login_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Create Classroom
    class_resp = client.post("/api/classrooms", json={
        "subject": "Document Processing CS501",
        "description": "Classroom for document processing tests"
    }, headers=headers)
    assert class_resp.status_code == 201
    classroom_id = class_resp.json()["id"]

    try:
        # --- TEST 1: PDF Extraction & Persistence ---
        pdf_path = f"temp_test_{timestamp}.pdf"
        pdf_text = "Normalization is a database design technique that reduces data redundancy."
        create_sample_pdf(pdf_path, pdf_text)
        temp_files.append(pdf_path)

        with open(pdf_path, "rb") as f:
            pdf_upload = client.post("/api/resources/upload", data={
                "classroom_id": classroom_id,
                "title": "Database Normalization PDF",
                "description": "Syllabus on normalization",
                "tags": "db,pdf"
            }, files={"file": ("normalization.pdf", f, "application/pdf")}, headers=headers)

        assert pdf_upload.status_code == 201, f"PDF Upload failed: {pdf_upload.text}"
        pdf_data = pdf_upload.json()
        pdf_id = pdf_data["id"]
        assert pdf_data["extraction_status"] == "PENDING"
        print(f"[PASS 1] PDF Uploaded with status PENDING. Resource ID: {pdf_id}")

        # Process background task explicitly for test verification
        process_uploaded_resource(pdf_id)

        status_resp = client.get(f"/api/resources/{pdf_id}/extraction-status", headers=headers)
        assert status_resp.status_code == 200
        assert status_resp.json()["status"] == "COMPLETED"
        print(f"[PASS 2] GET /api/resources/{pdf_id}/extraction-status returned COMPLETED")

        # Verify DB extracted text
        res_details = client.get(f"/api/resources/{pdf_id}", headers=headers).json()
        assert "Normalization" in res_details["extracted_text"]
        assert res_details["extraction_error"] is None
        print(f"[PASS 3] Extracted text persisted in DB: '{res_details['extracted_text'][:40]}...'")

        # --- TEST 2: DOCX Extraction ---
        docx_path = f"temp_test_{timestamp}.docx"
        docx_text = "Software Architecture Principles and Patterns"
        create_sample_docx(docx_path, docx_text)
        temp_files.append(docx_path)

        with open(docx_path, "rb") as f:
            docx_upload = client.post("/api/resources/upload", data={
                "classroom_id": classroom_id,
                "title": "Architecture Principles DOCX",
            }, files={"file": ("architecture.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=headers)
        
        assert docx_upload.status_code == 201
        docx_id = docx_upload.json()["id"]
        process_uploaded_resource(docx_id)

        docx_details = client.get(f"/api/resources/{docx_id}", headers=headers).json()
        assert docx_details["extraction_status"] == "COMPLETED"
        assert "Software Architecture" in docx_details["extracted_text"]
        print(f"[PASS 4] DOCX Text Extracted successfully: '{docx_details['extracted_text']}'")

        # --- TEST 3: PPTX Extraction ---
        pptx_path = f"temp_test_{timestamp}.pptx"
        pptx_text = "Machine Learning Foundations Lecture 1"
        create_sample_pptx(pptx_path, pptx_text)
        temp_files.append(pptx_path)

        with open(pptx_path, "rb") as f:
            pptx_upload = client.post("/api/resources/upload", data={
                "classroom_id": classroom_id,
                "title": "ML Lecture Slides PPTX",
            }, files={"file": ("ml_intro.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, headers=headers)
        
        assert pptx_upload.status_code == 201
        pptx_id = pptx_upload.json()["id"]
        process_uploaded_resource(pptx_id)

        pptx_details = client.get(f"/api/resources/{pptx_id}", headers=headers).json()
        assert pptx_details["extraction_status"] == "COMPLETED"
        assert "Machine Learning" in pptx_details["extracted_text"]
        print(f"[PASS 5] PPTX Text Extracted successfully: '{pptx_details['extracted_text']}'")

        # --- TEST 4: TXT Extraction ---
        txt_path = f"temp_test_{timestamp}.txt"
        txt_text = "Distributed Systems Consensus Algorithms Paxos Raft"
        create_sample_txt(txt_path, txt_text)
        temp_files.append(txt_path)

        with open(txt_path, "rb") as f:
            txt_upload = client.post("/api/resources/upload", data={
                "classroom_id": classroom_id,
                "title": "Consensus Algorithms TXT",
            }, files={"file": ("consensus.txt", f, "text/plain")}, headers=headers)
        
        assert txt_upload.status_code == 201
        txt_id = txt_upload.json()["id"]
        process_uploaded_resource(txt_id)

        txt_details = client.get(f"/api/resources/{txt_id}", headers=headers).json()
        assert txt_details["extraction_status"] == "COMPLETED"
        assert "Distributed Systems" in txt_details["extracted_text"]
        print(f"[PASS 6] TXT Text Extracted successfully: '{txt_details['extracted_text']}'")

        # --- TEST 5: Workspace Upload & Extraction Status API ---
        with open(pdf_path, "rb") as f:
            ws_upload = client.post("/api/workspace/upload", data={
                "title": "Private Normalization Note File"
            }, files={"file": ("private_norm.pdf", f, "application/pdf")}, headers=headers)
        
        assert ws_upload.status_code == 201
        ws_id = ws_upload.json()["id"]
        process_workspace_resource(ws_id)

        ws_status_resp = client.get(f"/api/workspace/{ws_id}/extraction-status", headers=headers)
        assert ws_status_resp.status_code == 200
        assert ws_status_resp.json()["status"] == "COMPLETED"

        ws_details = client.get(f"/api/workspace/{ws_id}", headers=headers).json()
        assert "Normalization" in ws_details["extracted_text"]
        print(f"[PASS 7] Workspace Resource extraction status API & DB persistence verified")

        # --- TEST 6: Search Readiness using ILIKE ---
        db = SessionLocal()
        try:
            current_user = db.query(User).filter(User.email == fac_email).first()
            search_results = search_resources_by_text(db, "normalization", current_user)
            assert len(search_results) >= 1
            assert str(search_results[0].id) == pdf_id
            print(f"[PASS 8] PostgreSQL ILIKE search readiness verified. Found resource: {search_results[0].title}")

            ws_search_results = search_workspace_files_by_text(db, "normalization", current_user)
            assert len(ws_search_results) >= 1
            assert str(ws_search_results[0].id) == ws_id
            print(f"[PASS 9] Workspace PostgreSQL ILIKE search readiness verified. Found file: {ws_search_results[0].title}")
        finally:
            db.close()

        # Verify metadata expansion
        assert res_details["preview_available"] is True
        assert res_details["extracted_text_available"] is True
        assert res_details["processing_status"] == "COMPLETED"
        assert res_details["ocr_status"] == "NOT_APPLICABLE"

        # --- TEST 7: Reprocessing API Endpoints ---
        reproc_resp = client.post(f"/api/resources/{pdf_id}/reprocess", headers=headers)
        assert reproc_resp.status_code == 200
        assert reproc_resp.json()["extraction_status"] == "COMPLETED"
        print(f"[PASS 10] POST /api/resources/{pdf_id}/reprocess returned COMPLETED")

        ws_reproc_resp = client.post(f"/api/workspace/{ws_id}/reprocess", headers=headers)
        assert ws_reproc_resp.status_code == 200
        assert ws_reproc_resp.json()["extraction_status"] == "COMPLETED"
        print(f"[PASS 11] POST /api/workspace/{ws_id}/reprocess returned COMPLETED")

        # --- TEST 8: Image / Missing OCR Graceful Handling ---
        img_upload = client.post("/api/resources/upload", data={
            "classroom_id": classroom_id,
            "title": "Diagram Image",
        }, files={"file": ("chart.png", b"\x89PNG\r\n\x1a\nDummy Image Content", "image/png")}, headers=headers)
        assert img_upload.status_code == 201
        img_id = img_upload.json()["id"]
        
        # Process image
        process_uploaded_resource(img_id)
        img_details = client.get(f"/api/resources/{img_id}", headers=headers).json()
        assert img_details["extraction_status"] in ("COMPLETED", "FAILED")
        assert "ocr_status" in img_details
        if img_details["extraction_status"] == "FAILED":
            assert img_details["extraction_error"] is not None
            print(f"[PASS 12] OCR Graceful Failure handling verified. OCR Status: '{img_details['ocr_status']}' | Error: '{img_details['extraction_error']}'")
        else:
            print("[PASS 12] Image OCR completed successfully")

    finally:
        # Cleanup temporary test files
        for tf in temp_files:
            if os.path.exists(tf):
                try:
                    os.remove(tf)
                except Exception:
                    pass
        # Clean test classroom
        client.delete(f"/api/classrooms/{classroom_id}", headers=headers)

    print("============================================================")
    print("ALL DOCUMENT PROCESSING & TEXT EXTRACTION TESTS PASSED 100%")
    print("============================================================\n")

if __name__ == "__main__":
    from cleanup_test_data import run_cleanup
    try:
        test_document_processing_suite()
    finally:
        print("\n[TEARDOWN] Cleaning up document processing test data...")
        run_cleanup(dry_run=False)

